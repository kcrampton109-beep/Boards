"""Tuesday weekly run: refresh everything, grade last week, price the new week, write the Excel workbook and the recap.
Usage: python weekly.py <staging dir>
Last output line is JSON: message, report (list of lines), workbook, page, state."""
import datetime as dt, json, pathlib, shutil, sys
import pandas as pd

ROOT = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
import engine

EXPORTS = ROOT / "exports"
EXPORTS.mkdir(exist_ok=True)


def rec(es):
    w = sum(e["status"] == "W" for e in es)
    l = sum(e["status"] == "L" for e in es)
    p = sum(e["status"] == "P" for e in es)
    u = sum(e.get("units") or 0 for e in es)
    return f"{w}-{l}" + (f"-{p}" if p else ""), u


def workbook(state, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="13233A")

    def sheet(name, cols, rows, widths=None):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, max(10, min(40, len(c) + 4)))
        return ws

    wb.remove(wb.active)
    games = {g["game_id"]: g for g in state["games"]}
    sheet("Plays", ["Tier", "Game", "Kickoff", "Pick", "Odds", "Book", "Edge", "Playable to", "Cover %", "Note"],
          [[p["tier"], games[p["game_id"]]["matchup"], games[p["game_id"]]["kickoff"], p["pick"],
            p["odds"] if p["market"] != "teaser" else "teaser", p["book"], p.get("edge"), p.get("limit"), p.get("cover"), p.get("note")]
           for p in state["plays"]], {"Game": 14, "Kickoff": 24, "Pick": 34, "Note": 60})
    sheet("Games", ["Week", "Away", "Home", "Kickoff", "Model line (home)", "Consensus", "DK", "FD", "MGM", "CBS open", "CBS now",
                    "Home bets %", "Home money %", "Model total", "Total", "Over bets %", "Wind mph", "Away QB", "Home QB"],
          [[g["week"], g["away"], g["home"], g["kickoff"], g["model_line"], g["cons_spread"],
            (g["books"].get("DK") or {}).get("sp_home"), (g["books"].get("FD") or {}).get("sp_home"), (g["books"].get("MGM") or {}).get("sp_home"),
            (g.get("cbs") or {}).get("open_spread_home"), (g.get("cbs") or {}).get("spread_home"), g["tix_home"], g["money_home"],
            g["model_total"], g["cons_total"], g["tix_over"], (g.get("wx") or {}).get("wind"),
            (g["qb"].get("away") or {}).get("expected"), (g["qb"].get("home") or {}).get("expected")] for g in state["games"]],
          {"Kickoff": 24})
    sheet("Bet Log", ["Logged", "Week", "Tier", "Game", "Pick", "Odds", "Book", "Stake", "Result", "Units", "Close", "CLV", "Final"],
          [[e["flagged"][:16].replace("T", " "), e["week"], e["tier"], e["matchup"], e["pick"], e["odds"], e["book"], e["stake"],
            e["status"], e.get("units"), e.get("close"), e.get("clv"), e.get("final")] for e in state["ledger"]],
          {"Logged": 18, "Pick": 34, "Final": 22})
    sm = state["summary"]
    sheet("Records", ["Tier", "W", "L", "P", "Pending", "Units", "ROI", "Avg CLV", "Beat close"],
          [[t, d["w"], d["l"], d["p"], d["open"], d["units"], d["roi"], d["clv_avg"], d["beat_close"]] for t, d in sm.items()])
    sheet("Ratings", ["Rank", "Team", "Rating", "Market", "Points", "EPA/play net", "QB", "QB EPA/db"],
          [[r["rank"], r["team"], r["net"], r["mkt"], r["pts"], r["epa"], r["qb"], r["qb_val"]] for r in state["ratings"]], {"QB": 22})
    tdrows = []
    for gm in state["td"]:
        for t in gm["teams"]:
            for p in t["players"]:
                tdrows.append([gm["matchup"], t["team"], p["player"], p["pos"], p["prob"], p["dk"], p["fd"], p["book_prob"], "VALUE" if p["value"] else ""])
    sheet("TD sheet", ["Game", "Team", "Player", "Pos", "Model chance", "DK", "FD", "Book chance", "Value"], tdrows, {"Player": 24})
    sheet("Injuries", ["Team", "Player", "Pos", "Status", "Injury", "Latest"],
          [[x["team"], x["player"], x["pos"], x["status"], x["detail"], x["comment"]] for x in state["injuries"]], {"Latest": 80, "Injury": 30})
    wb.save(path)


def main(dest):
    engine.main(["--refresh"])
    state = json.loads(engine.STATE.read_text())
    week = state["week"]
    led = state["ledger"]
    last = [e for e in led if e["week"] == week - 1 and e["status"] != "open"]
    lines = [f"NFL week {week - 1} paper results:"]
    for tier in ("STRONG", "LEAN", "WEATHER", "TEASER", "RLM", "TREND", "TD"):
        es = [e for e in last if e["tier"] == tier]
        if es:
            r, u = rec(es)
            lines.append(f"  {tier.title()}: {r}" + ("" if tier == "TEASER" else f", {u:+.2f}u"))
    if len(lines) == 1:
        lines.append("  no graded plays")
    sm = state["summary"]["STRONG"]
    lines.append(f"Season, strong spreads: {sm['w']}-{sm['l']}, {sm['units']:+.2f}u, CLV {sm['clv_avg'] if sm['clv_avg'] is not None else '–'} pts")
    cur = [p for p in state["plays"] if p.get("week") == week + 1 or p.get("week") == week]
    for tier in ("STRONG", "WEATHER", "TEASER"):
        ps = [p for p in cur if p["tier"] == tier]
        if ps:
            lines.append(f"{tier.title()} now: " + "; ".join(f"{p['pick']} ({p['matchup']})" for p in ps[:6]))
    rl = [p for p in cur if p["tier"] == "RLM"]
    if rl:
        lines.append(f"RLM flags (tracked): {len(rl)}")
    tr = [p for p in cur if p["tier"] == "TREND"]
    for p in tr[:4]:
        lines.append(f"Trend: {p['pick']} - {p['note']}")
    qb = [f"{t}: {q['note']}" for t, q in state["qbs"].items() if q.get("note")]
    if qb:
        lines.append("QB changes: " + "; ".join(qb))
    stamp = dt.date.today().isoformat()
    wbp = EXPORTS / f"NFL_Week{week + 1 if all(g['started'] for g in state['games'] if g['week'] == week) else week}_{stamp}.xlsx"
    workbook(state, wbp)
    dest = pathlib.Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    shutil.copy(ROOT / "page" / "index.html", dest / "index.html")
    shutil.copy(engine.STATE, dest / "state.json")
    msg = f"NFL weekly: week {week - 1} graded, new week priced. Workbook attached."
    print(json.dumps(dict(message=msg, report=lines, workbook=str(wbp.resolve()), page=str((dest / "index.html").resolve()),
                          state=str((dest / "state.json").resolve()))))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "stage"))
